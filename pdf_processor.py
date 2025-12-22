#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
PDF处理模块
将PDF文件转换为DOCX格式，以便使用现有的文档处理流程
"""

import os
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Pt
from io import BytesIO

# 导入日志工具
try:
    from utils import setup_logger, log_info, log_warning, log_error
    from parallel_processor import _logger
except Exception:
    _logger = None


def pdf_to_docx(pdf_path, docx_path=None, debug_mode=False):
    """
    将PDF文件转换为DOCX格式
    
    Args:
        pdf_path: PDF文件路径
        docx_path: 输出的DOCX文件路径，如果为None则自动生成
        debug_mode: 是否显示调试信息
    
    Returns:
        转换后的DOCX文件路径，失败返回None
    """
    try:
        # 若PDF为纯图片（无文本），优先使用 minerU 转换
        if _is_image_only_pdf(pdf_path):
            log_info(_logger, f"检测到图片型PDF，调用 minerU 转换: {pdf_path}")
            docx_via_mineru = _convert_with_mineru(pdf_path, docx_path)
            if docx_via_mineru:
                return docx_via_mineru
            else:
                log_warning(_logger, "minerU 转换失败，回退至内置转换流程")

        # 尝试使用PyMuPDF (fitz) - 更强大的PDF处理库
        try:
            import fitz  # PyMuPDF
            return _pdf_to_docx_pymupdf(pdf_path, docx_path, debug_mode)
        except ImportError:
            pass
        
        # 尝试使用pdfplumber - 另一个PDF处理库
        try:
            import pdfplumber
            return _pdf_to_docx_pdfplumber(pdf_path, docx_path, debug_mode)
        except ImportError:
            pass
        
        # 尝试使用PyPDF2 - 基础PDF处理库
        try:
            import PyPDF2
            return _pdf_to_docx_pypdf2(pdf_path, docx_path, debug_mode)
        except ImportError:
            pass
        
        raise ImportError("未找到PDF处理库，请安装 PyMuPDF、pdfplumber 或 PyPDF2")
        
    except Exception as e:
        log_error(_logger, pdf_path, e, type(e).__name__, None)
        return None


def _is_image_only_pdf(pdf_path: str | Path) -> bool:
    """快速检测PDF是否无可提取文本（视为图片型PDF）"""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return False

    try:
        doc = fitz.open(pdf_path)
        text_len = 0
        for page in doc:
            text_len += len(page.get_text().strip())
            if text_len > 0:
                break
        doc.close()
        return text_len == 0
    except Exception:
        return False


def _convert_with_mineru(pdf_path: str | Path, docx_path: str | Path | None):
    """
    使用 minerU 将图片型 PDF 转换为 DOCX。
    需要环境变量 MINERU_CMD 指定命令模板，默认: "mineru convert --input {input} --output {output}"
    """
    cmd_tpl = os.getenv("MINERU_CMD", "mineru convert --input {input} --output {output}")
    if not cmd_tpl:
        log_warning(_logger, "未配置 MINERU_CMD，无法调用 minerU 进行图片型PDF转换")
        return None

    if docx_path is None:
        docx_path = str(Path(pdf_path).with_suffix(".docx"))

    try:
        cmd = cmd_tpl.format(input=str(pdf_path), output=str(docx_path))
        log_info(_logger, f"执行 minerU 命令: {cmd}")
        result = subprocess.run(cmd, shell=True, capture_output=True, text=True, timeout=300)
        if result.returncode != 0:
            log_error(_logger, pdf_path, f"minerU 转换失败: {result.stderr}", "MinerUError", result.stderr)
            return None
        if not Path(docx_path).exists():
            log_error(_logger, pdf_path, "minerU 转换后未找到输出文件", "MinerUError", None)
            return None
        return docx_path
    except Exception as e:
        log_error(_logger, pdf_path, e, type(e).__name__, None)
        return None


def _pdf_to_docx_pymupdf(pdf_path, docx_path=None, debug_mode=False):
    """使用PyMuPDF (fitz) 转换PDF到DOCX"""
    import fitz  # PyMuPDF
    
    if docx_path is None:
        docx_path = str(Path(pdf_path).with_suffix('.docx'))
    
    doc = Document()
    
    try:
        pdf_doc = fitz.open(pdf_path)
        
        log_info(_logger, f"  正在处理PDF: {pdf_path} (共 {len(pdf_doc)} 页)")
        
        for page_num in range(len(pdf_doc)):
            page = pdf_doc[page_num]
            
            # 提取文本
            text = page.get_text()
            
            if text.strip():
                # 按段落分割文本
                paragraphs = text.split('\n\n')
                for para_text in paragraphs:
                    para_text = para_text.strip()
                    if para_text:
                        # 检查是否是标题（简单启发式：短文本且可能包含数字）
                        if len(para_text) < 100 and (para_text.isupper() or para_text[0].isdigit()):
                            para = doc.add_heading(para_text, level=2)
                        else:
                            para = doc.add_paragraph(para_text)
                        para.style.font.size = Pt(12)
            
            # 如果不是最后一页，添加分页符
            if page_num < len(pdf_doc) - 1:
                doc.add_page_break()
        
        pdf_doc.close()
        
        # 保存DOCX文件
        doc.save(docx_path)
        
        log_info(_logger, f"  PDF转换完成: {docx_path}")
        
        return docx_path
        
    except Exception as e:
        log_error(_logger, pdf_path, e, type(e).__name__, None)
        raise


def _pdf_to_docx_pdfplumber(pdf_path, docx_path=None, debug_mode=False):
    """使用pdfplumber转换PDF到DOCX"""
    import pdfplumber
    
    if docx_path is None:
        docx_path = str(Path(pdf_path).with_suffix('.docx'))
    
    doc = Document()
    
    try:
        with pdfplumber.open(pdf_path) as pdf:
            log_info(_logger, f"  正在处理PDF: {pdf_path} (共 {len(pdf.pages)} 页)")
            
            for page_num, page in enumerate(pdf.pages):
                # 提取文本
                text = page.extract_text()
                
                if text:
                    # 按段落分割文本
                    paragraphs = text.split('\n\n')
                    for para_text in paragraphs:
                        para_text = para_text.strip()
                        if para_text:
                            # 检查是否是标题
                            if len(para_text) < 100 and (para_text.isupper() or para_text[0].isdigit()):
                                para = doc.add_heading(para_text, level=2)
                            else:
                                para = doc.add_paragraph(para_text)
                            para.style.font.size = Pt(12)
                
                # 如果不是最后一页，添加分页符
                if page_num < len(pdf.pages) - 1:
                    doc.add_page_break()
        
        # 保存DOCX文件
        doc.save(docx_path)
        
        log_info(_logger, f"  PDF转换完成: {docx_path}")
        
        return docx_path
        
    except Exception as e:
        log_error(_logger, pdf_path, e, type(e).__name__, None)
        raise


def _pdf_to_docx_pypdf2(pdf_path, docx_path=None, debug_mode=False):
    """使用PyPDF2转换PDF到DOCX（基础功能）"""
    import PyPDF2
    
    if docx_path is None:
        docx_path = str(Path(pdf_path).with_suffix('.docx'))
    
    doc = Document()
    
    try:
        with open(pdf_path, 'rb') as pdf_file:
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            log_info(_logger, f"  正在处理PDF: {pdf_path} (共 {len(pdf_reader.pages)} 页)")
            
            for page_num, page in enumerate(pdf_reader.pages):
                # 提取文本
                text = page.extract_text()
                
                if text:
                    # 按段落分割文本
                    paragraphs = text.split('\n\n')
                    for para_text in paragraphs:
                        para_text = para_text.strip()
                        if para_text:
                            # 检查是否是标题
                            if len(para_text) < 100 and (para_text.isupper() or para_text[0].isdigit()):
                                para = doc.add_heading(para_text, level=2)
                            else:
                                para = doc.add_paragraph(para_text)
                            para.style.font.size = Pt(12)
                
                # 如果不是最后一页，添加分页符
                if page_num < len(pdf_reader.pages) - 1:
                    doc.add_page_break()
        
        # 保存DOCX文件
        doc.save(docx_path)
        
        log_info(_logger, f"  PDF转换完成: {docx_path}")
        
        return docx_path
        
    except Exception as e:
        log_error(_logger, pdf_path, e, type(e).__name__, None)
        raise


def is_pdf_file(file_path):
    """检查文件是否是PDF格式"""
    if isinstance(file_path, (str, Path)):
        path = Path(file_path)
        return path.suffix.lower() == '.pdf'
    return False

